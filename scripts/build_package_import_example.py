"""Build the manager Excel example from an existing package JSON."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "content" / "packages" / "3-day-classic-experience-in-southern-laos.json"
REVIEWED = ROOT / "content" / "translations" / "reviewed_en_fr_th.json"
OUT = ROOT / "wp-content" / "plugins" / "watphou-core" / "data" / "watphou-package-import-example.xlsx"
EXAMPLE_SLUG = "3-day-classic-experience-in-southern-laos"

TOUR_HEADERS = [
    (
        "slug",
        "SLUG = web address name of the tour. Use lowercase letters and hyphens, "
        "for example 3-day-classic-experience-in-southern-laos. Keep the same slug to UPDATE an existing tour. "
        "Use a new slug to ADD a tour. Do not put spaces or accented letters in the slug.",
    ),
    (
        "title",
        "TITLE = the public name visitors see as the main heading, for example 3-Day Classic Experience in Southern Laos.",
    ),
    (
        "headline",
        "HEADLINE = one short sales line under the title (the promise of the tour).",
    ),
    (
        "duration",
        "DURATION = the length visitors see, for example 3 days / 2 nights or 1 day.",
    ),
    (
        "duration_term",
        "DURATION_TERM = which menu the tour belongs to. Type exactly one of: 1-day, 2-day, 3-day, 4-6-day.",
    ),
    (
        "destinations",
        "DESTINATIONS = places this tour visits. Comma-separated list using only: bolaven-plateau, 4000-islands, champasak, pakse.",
    ),
    (
        "price_from",
        "PRICE_FROM = starting price in US dollars. Type a number when the real price is confirmed, or XX until then. Do not invent a price.",
    ),
    (
        "bestseller",
        "BESTSELLER = yes to show the tour on the homepage strip, no to hide it from that strip.",
    ),
    (
        "priority",
        "PRIORITY = homepage order. A higher number appears first. 0 is fine if you are not sure.",
    ),
    (
        "code",
        "CODE = internal package number from the Watphou catalogue, for example 3.1 or 1.1.",
    ),
    (
        "dream",
        "DREAM = the opening paragraph that describes the feeling of the tour (two to five sentences).",
    ),
    (
        "highlights",
        "HIGHLIGHTS = the bullet list of the best moments. Put one highlight per line (Alt+Enter inside the cell).",
    ),
    (
        "included",
        "INCLUDED = what the price covers. One item per line (vehicle, meals, tickets, and so on).",
    ),
    (
        "excluded",
        "EXCLUDED = what the price does not cover. One item per line.",
    ),
    (
        "upgrades",
        "UPGRADES = optional extras the visitor can add (guide, better hotel, zipline). One item per line.",
    ),
    (
        "cta",
        "CTA = Call To Action. The last sentence that invites the visitor to book, for example a WhatsApp request.",
    ),
]

DAY_HEADERS = [
    (
        "slug",
        "SLUG = must be the same web address name as on the Tours sheet, so WordPress knows which tour these days belong to.",
    ),
    (
        "day",
        "DAY = the day number in the itinerary: 1 for the first day, 2 for the second day, and so on.",
    ),
    (
        "title",
        "TITLE = short name of that day, for example Pakse to the 4,000 Islands.",
    ),
    (
        "body",
        "BODY = what happens that day (places, meals, overnight). Write in full sentences.",
    ),
]


def join_list(items: list[str]) -> str:
    return "\n".join(items)


def reviewed_package_lang(items: list[dict], slug: str, lang: str) -> dict:
    parts = [
        i
        for i in items
        if i.get("kind") == "package" and str(i.get("location") or "").endswith(slug)
    ]
    data: dict = {
        "highlights": {},
        "included": {},
        "excluded": {},
        "upgrades": {},
        "itinerary": {},
    }
    for row in parts:
        field = str(row.get("field") or "")
        text = str(row.get(lang) or "").strip()
        if field in {
            "title",
            "headline",
            "duration",
            "departure",
            "price_note",
            "dream",
            "cta",
            "code",
        }:
            data[field] = text
            continue
        listed = re.match(r"^(highlights|included|excluded|upgrades)\[(\d+)\]$", field)
        if listed:
            data[listed.group(1)][int(listed.group(2))] = text
            continue
        day_field = re.match(r"^itinerary\.day(\d+)\.(title|body)$", field)
        if day_field:
            day = int(day_field.group(1))
            data["itinerary"].setdefault(day, {"day": day, "title": "", "body": ""})
            data["itinerary"][day][day_field.group(2)] = text
    for key in ("highlights", "included", "excluded", "upgrades"):
        data[key] = [data[key][k] for k in sorted(data[key]) if str(data[key][k]).strip()]
    data["itinerary"] = [data["itinerary"][k] for k in sorted(data["itinerary"])]
    return data


def lang_sheet_headers(max_day: int) -> list[tuple[str, str]]:
    extra: list[tuple[str, str]] = []
    for n in range(1, max_day + 1):
        extra.append(
            (
                f"day{n}_title",
                f"DAY{n}_TITLE = short name of day {n} in this language. Keep the same meaning as the English Itinerary sheet.",
            )
        )
        extra.append(
            (
                f"day{n}_body",
                f"DAY{n}_BODY = what happens on day {n}, written in this language.",
            )
        )
    return TOUR_HEADERS + extra


def write_lang_sheet(
    wb: Workbook,
    title: str,
    tab_color: str,
    en: dict,
    lang_data: dict,
    header_fill: PatternFill,
    header_font: Font,
    wrap: Alignment,
) -> list[tuple[str, str]]:
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    max_day = max(len(en.get("itinerary") or []), len(lang_data.get("itinerary") or []), 1)
    headers = lang_sheet_headers(max_day)
    for col, (name, hint) in enumerate(headers, start=1):
        cell = ws.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.comment = Comment(hint, "Watphou Travels", 300, 100)
        cell.comment.width = "320pt"
        cell.comment.height = "110pt"
        wide = name in {"dream", "highlights", "included", "excluded", "upgrades"} or name.endswith("_body")
        ws.column_dimensions[get_column_letter(col)].width = 50 if wide else 28
    values = {
        "slug": en["slug"],
        "title": lang_data.get("title", ""),
        "headline": lang_data.get("headline", ""),
        "duration": lang_data.get("duration", ""),
        "duration_term": en.get("duration_term", ""),
        "destinations": ", ".join(en.get("destinations", [])),
        "price_from": en.get("price_from", "XX"),
        "bestseller": "yes" if en.get("bestseller") else "no",
        "priority": en.get("priority", 0),
        "code": lang_data.get("code") or en.get("code", ""),
        "dream": lang_data.get("dream", ""),
        "highlights": join_list(lang_data.get("highlights", [])),
        "included": join_list(lang_data.get("included", [])),
        "excluded": join_list(lang_data.get("excluded", [])),
        "upgrades": join_list(lang_data.get("upgrades", [])),
        "cta": lang_data.get("cta", ""),
    }
    for day in lang_data.get("itinerary") or []:
        n = int(day.get("day") or 0)
        if n < 1:
            continue
        values[f"day{n}_title"] = day.get("title", "")
        values[f"day{n}_body"] = day.get("body", "")
    for col, (name, _hint) in enumerate(headers, start=1):
        cell = ws.cell(2, col, values.get(name, ""))
        cell.alignment = wrap
    ws.row_dimensions[2].height = 110
    ws.freeze_panes = "A2"
    return headers


def main() -> None:
    data = json.loads(SRC.read_text(encoding="utf-8"))
    OUT.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    how = wb.active
    how.title = "How to use"
    tours = wb.create_sheet("Tours")
    days = wb.create_sheet("Itinerary")

    how["A1"] = "How to add or update a Watphou Travels tour"
    how["A1"].font = Font(bold=True, size=14)
    instructions = [
        "",
        "1. This file is an example. The Tours sheet already contains the live 3-Day Classic Experience package so you can copy the layout.",
        "2. To UPDATE that tour: keep the same slug, change the text, save, then upload the file in WordPress (Watphou → Add tours Excel).",
        "3. To ADD a new tour: copy the example row, change the slug (new web address name) and the title, then fill the other columns.",
        "4. Put each day on the Itinerary sheet. Use the same slug as the Tours row.",
        "5. Pictures are not in this file. After upload, open Watphou → Quick edit tours → Change photo.",
        "6. Leave price_from as XX until the real price is confirmed.",
        "7. Save as .xlsx (Excel workbook). Do not use .xls or Google-Sheets-only formats.",
        "",
        "Sheet FR = French tour text. Sheet TH = Thai tour text. Keep the same slug as on Tours.",
        "On FR and TH, day titles and bodies are extra columns: day1_title, day1_body, day2_title, day2_body, …",
        "You can edit those sheets and upload the same file; WordPress will update /fr/ and /th/ for that tour.",
        "",
        "Allowed destinations: bolaven-plateau, 4000-islands, champasak, pakse",
        "Allowed duration_term: 1-day, 2-day, 3-day, 4-6-day",
        "",
        "Hover the yellow header cells (small red triangle) to read what each column means.",
        "There is also a sheet named Column meanings with the same explanations.",
        "",
    ]
    for i, line in enumerate(instructions, start=2):
        how[f"A{i}"] = line
        how[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    how.column_dimensions["A"].width = 110

    header_fill = PatternFill("solid", fgColor="F4D35E")
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, (name, hint) in enumerate(TOUR_HEADERS, start=1):
        cell = tours.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.comment = Comment(hint, "Watphou Travels", 300, 100)
        cell.comment.width = "320pt"
        cell.comment.height = "110pt"
        tours.column_dimensions[get_column_letter(col)].width = 28 if name != "dream" else 50

    tour_values = {
        "slug": data["slug"],
        "title": data["title"],
        "headline": data.get("headline", ""),
        "duration": data.get("duration", ""),
        "duration_term": data.get("duration_term", ""),
        "destinations": ", ".join(data.get("destinations", [])),
        "price_from": data.get("price_from", "XX"),
        "bestseller": "yes" if data.get("bestseller") else "no",
        "priority": data.get("priority", 0),
        "code": data.get("code", ""),
        "dream": data.get("dream", ""),
        "highlights": join_list(data.get("highlights", [])),
        "included": join_list(data.get("included", [])),
        "excluded": join_list(data.get("excluded", [])),
        "upgrades": join_list(data.get("upgrades", [])),
        "cta": data.get("cta", ""),
    }
    for col, (name, _hint) in enumerate(TOUR_HEADERS, start=1):
        cell = tours.cell(2, col, tour_values[name])
        cell.alignment = wrap
    tours.row_dimensions[2].height = 90
    tours.freeze_panes = "A2"
    tours.auto_filter.ref = f"A1:{get_column_letter(len(TOUR_HEADERS))}2"

    dur_dv = DataValidation(type="list", formula1='"1-day,2-day,3-day,4-6-day"', allow_blank=True)
    yes_dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    tours.add_data_validation(dur_dv)
    tours.add_data_validation(yes_dv)
    dur_dv.add("E2:E200")
    yes_dv.add("H2:H200")

    for col, (name, hint) in enumerate(DAY_HEADERS, start=1):
        cell = days.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.comment = Comment(hint, "Watphou Travels", 300, 100)
        cell.comment.width = "320pt"
        cell.comment.height = "110pt"
        days.column_dimensions[get_column_letter(col)].width = 36 if name != "body" else 70

    for i, day in enumerate(data.get("itinerary", []), start=2):
        days.cell(i, 1, data["slug"]).alignment = wrap
        days.cell(i, 2, int(day.get("day", i - 1)))
        days.cell(i, 3, day.get("title", "")).alignment = wrap
        days.cell(i, 4, day.get("body", "")).alignment = wrap
        days.row_dimensions[i].height = 45
    days.freeze_panes = "A2"

    reviewed_items = []
    if REVIEWED.is_file():
        reviewed_items = json.loads(REVIEWED.read_text(encoding="utf-8"))
    fr_data = reviewed_package_lang(reviewed_items, EXAMPLE_SLUG, "fr")
    th_data = reviewed_package_lang(reviewed_items, EXAMPLE_SLUG, "th")
    fr_headers = write_lang_sheet(wb, "FR", "2E86AB", data, fr_data, header_fill, header_font, wrap)
    th_headers = write_lang_sheet(wb, "TH", "C73E1D", data, th_data, header_fill, header_font, wrap)

    meanings = wb.create_sheet("Column meanings")
    meanings["A1"] = "Column"
    meanings["B1"] = "What it means"
    meanings["A1"].font = header_font
    meanings["B1"].font = header_font
    meanings["A1"].fill = header_fill
    meanings["B1"].fill = header_fill
    row_i = 2
    meanings.cell(row_i, 1, "Tours sheet")
    meanings.cell(row_i, 1).font = Font(bold=True)
    row_i += 1
    for name, hint in TOUR_HEADERS:
        meanings.cell(row_i, 1, name)
        meanings.cell(row_i, 2, hint).alignment = Alignment(wrap_text=True, vertical="top")
        meanings.row_dimensions[row_i].height = 48
        row_i += 1
    row_i += 1
    meanings.cell(row_i, 1, "Itinerary sheet")
    meanings.cell(row_i, 1).font = Font(bold=True)
    row_i += 1
    for name, hint in DAY_HEADERS:
        meanings.cell(row_i, 1, name)
        meanings.cell(row_i, 2, hint).alignment = Alignment(wrap_text=True, vertical="top")
        meanings.row_dimensions[row_i].height = 48
        row_i += 1
    row_i += 1
    meanings.cell(row_i, 1, "FR sheet (French) and TH sheet (Thai)")
    meanings.cell(row_i, 1).font = Font(bold=True)
    row_i += 1
    meanings.cell(row_i, 1, "same slug")
    meanings.cell(row_i, 2, "Must match the English Tours sheet so WordPress updates the French or Thai copy of that tour.").alignment = Alignment(
        wrap_text=True, vertical="top"
    )
    meanings.row_dimensions[row_i].height = 48
    row_i += 1
    for name, hint in fr_headers:
        if name in {h[0] for h in TOUR_HEADERS}:
            continue
        meanings.cell(row_i, 1, name)
        meanings.cell(row_i, 2, hint).alignment = Alignment(wrap_text=True, vertical="top")
        meanings.row_dimensions[row_i].height = 48
        row_i += 1
    meanings.column_dimensions["A"].width = 18
    meanings.column_dimensions["B"].width = 100

    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
