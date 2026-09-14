"""Dump reviewed translation Excel into JSON and print a summary."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "translations" / "Watphou_EN_FR_TH_reviewed.xlsx"
OUT = ROOT / "content" / "translations" / "reviewed_en_fr_th.json"


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Texts"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h or "").strip() for h in rows[0]]
    items = []
    for raw in rows[1:]:
        if not raw or not raw[0]:
            continue
        rec = {headers[i]: raw[i] for i in range(len(headers))}
        items.append(
            {
                "id": rec.get("ID"),
                "kind": str(rec.get("Type") or "").strip(),
                "location": str(rec.get("Location") or "").strip(),
                "field": str(rec.get("Field") or "").strip(),
                "en": str(rec.get("English (source)") or "").strip(),
                "fr": str(rec.get("French (reviewed draft)") or "").strip(),
                "th": str(rec.get("Thai (reviewed draft)") or "").strip(),
                "notes": str(rec.get("Reviewer notes") or "").strip(),
                "status": str(rec.get("Status") or "").strip(),
            }
        )

    kinds = Counter(i["kind"] for i in items)
    fields = Counter(i["field"] for i in items)
    locs = Counter(f"{i['kind']}:{i['location']}" for i in items)
    status = Counter(i["status"] for i in items)
    empty_fr = sum(1 for i in items if not i["fr"])
    empty_th = sum(1 for i in items if not i["th"])
    same_fr = sum(1 for i in items if i["fr"] and i["fr"] == i["en"])
    same_th = sum(1 for i in items if i["th"] and i["th"] == i["en"])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    plugin_copy = ROOT / "wp-content" / "plugins" / "watphou-core" / "data" / "reviewed_en_fr_th.json"
    plugin_copy.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")

    summary = {
        "rows": len(items),
        "kinds": dict(kinds),
        "fields": dict(fields),
        "status": dict(status),
        "empty_fr": empty_fr,
        "empty_th": empty_th,
        "same_as_en_fr": same_fr,
        "same_as_en_th": same_th,
        "locations": dict(locs),
        "plugin_copy": str(plugin_copy),
    }
    Path(ROOT / "content" / "translations" / "reviewed_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
